#!/usr/bin/env python3
"""Build a downloadable .docx from the live print page."""

from __future__ import annotations

import io
import re
import sys
import urllib.request
from pathlib import Path

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "public"
OUT = PUBLIC / "downloads" / "relay-life-test-report.docx"
FIGURES = PUBLIC / "figures"
PRINT_URL = "http://127.0.0.1:43217/print"

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_BODY = "Times New Roman"
EN_CODE = "Consolas"


def set_run_font(run, *, east=CN_BODY, ascii_font=EN_BODY, size=12, bold=False, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = ascii_font
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east)


def set_paragraph_format(p, *, first_indent=None, space_before=0, space_after=6, line=1.5, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if align is not None:
        p.alignment = align


def classes(tag: Tag) -> str:
    return " ".join(tag.get("class") or [])


def visible_text(node) -> str:
    return re.sub(r"\s+", " ", node.get_text(" ", strip=True)).strip()


def add_inline(p, node, *, bold=False, code=False, size=12):
    if isinstance(node, Comment):
        return
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        text = text.replace("\xa0", " ")
        if not text.strip() and text != " ":
            return
        run = p.add_run(text)
        if code:
            set_run_font(run, east=CN_BODY, ascii_font=EN_CODE, size=size - 1, bold=bold)
        else:
            set_run_font(run, east=CN_BODY, ascii_font=EN_BODY, size=size, bold=bold)
        return
    if not isinstance(node, Tag):
        return
    if node.name in {"script", "style"}:
        return
    next_bold = bold or node.name in {"strong", "b"}
    next_code = code or node.name == "code"
    if node.name == "br":
        p.add_run().add_break()
        return
    for child in node.children:
        add_inline(p, child, bold=next_bold, code=next_code, size=size)


class Builder:
    def __init__(self, html: str):
        self.soup = BeautifulSoup(html, "lxml")
        self.doc = Document()
        self._setup_document()
        self.pending_kicker = ""

    def _setup_document(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        normal = self.doc.styles["Normal"]
        normal.font.name = EN_BODY
        normal.font.size = Pt(12)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), EN_BODY)
        rfonts.set(qn("w:hAnsi"), EN_BODY)
        rfonts.set(qn("w:eastAsia"), CN_BODY)
        self.doc.core_properties.title = "继电器接点寿命测试软件 · 功能实现与关键技术分析"
        self.doc.core_properties.subject = "课题报告整理稿"
        self.doc.core_properties.language = "zh-CN"

    def add_heading_text(self, text: str, level: int):
        p = self.doc.add_heading(text, level=level)
        for run in p.runs:
            size = {0: 22, 1: 16, 2: 14, 3: 12}.get(level, 12)
            set_run_font(run, east=CN_HEAD, ascii_font=EN_BODY, size=size, bold=True)
        set_paragraph_format(p, first_indent=0, space_before=12 if level else 0, space_after=8, line=1.3)
        return p

    def add_body(self, node: Tag, *, indent=0.74, align=None, space_before=0, size=12):
        p = self.doc.add_paragraph()
        set_paragraph_format(p, first_indent=indent, space_before=space_before, space_after=6, align=align)
        add_inline(p, node, size=size)
        if not visible_text(node):
            return None
        return p

    def add_plain(self, text: str, *, bold=False, indent=0, align=None, size=12, space_before=0, space_after=6, color=None):
        p = self.doc.add_paragraph()
        set_paragraph_format(p, first_indent=indent, space_before=space_before, space_after=space_after, align=align)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
        return p

    def add_list(self, node: Tag, ordered: bool):
        for i, li in enumerate(node.find_all("li", recursive=False), start=1):
            prefix = f"{i}. " if ordered else "• "
            p = self.doc.add_paragraph()
            set_paragraph_format(p, first_indent=0, space_before=0, space_after=3, line=1.4)
            p.paragraph_format.left_indent = Cm(0.75)
            run = p.add_run(prefix)
            set_run_font(run, bold=False, size=12)
            for child in li.children:
                add_inline(p, child, size=12)

    def add_table(self, table_tag: Tag):
        rows = table_tag.find_all("tr")
        if not rows:
            return
        cols = max(len(tr.find_all(["th", "td"])) for tr in rows)
        if cols == 0:
            return
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, tr in enumerate(rows):
            cells = tr.find_all(["th", "td"])
            for c in range(cols):
                cell = table.cell(r, c)
                cell.text = ""
                p = cell.paragraphs[0]
                set_paragraph_format(p, first_indent=0, space_before=0, space_after=0, line=1.15)
                if c < len(cells):
                    is_head = cells[c].name == "th" or r == 0
                    add_inline(p, cells[c], bold=is_head, size=10)
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.font.size is None:
                            set_run_font(run, size=10)
        self.doc.add_paragraph()

    def add_image(self, src: str, caption: str):
        path = FIGURES / Path(src).name
        if not path.exists():
            self.add_plain(f"[缺图] {caption or src}", indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            return
        p = self.doc.add_paragraph()
        set_paragraph_format(p, first_indent=0, space_before=8, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run()
        run.add_picture(str(path), width=Cm(15.5))
        if caption:
            cap = self.doc.add_paragraph()
            set_paragraph_format(cap, first_indent=0, space_before=0, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            r = cap.add_run(caption)
            set_run_font(r, east=CN_HEAD, size=10.5, bold=True)

    def add_figure(self, fig: Tag):
        cap = fig.find("figcaption")
        caption = visible_text(cap) if cap else ""
        img = fig.find("img")
        if img and img.get("src"):
            self.add_image(img["src"], caption or visible_text(img))
            return
        if caption:
            self.add_plain(caption, bold=True, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_before=8)
        clone = BeautifulSoup(str(fig), "lxml")
        if clone.figcaption:
            clone.figcaption.decompose()
        bits = [t.strip() for t in clone.get_text("\n", strip=True).splitlines() if t.strip()]
        if bits:
            self.add_plain("；".join(bits), indent=0, size=10.5, space_after=8)

    def add_qa(self, article: Tag):
        q = article.find("p")
        if q:
            title = visible_text(q)
            p = self.doc.add_paragraph()
            set_paragraph_format(p, first_indent=0, space_before=8, space_after=4)
            r = p.add_run(title)
            set_run_font(r, bold=True, size=12, color=RGBColor(0x92, 0x3F, 0x0C))
        for para in article.find_all("p")[1:]:
            self.add_body(para, indent=0.0, size=11.5)

    def add_algo(self, article: Tag):
        header = article.find("header")
        if header:
            self.add_plain(visible_text(header), bold=True, indent=0, size=12, space_before=10, space_after=4, color=RGBColor(0x0F, 0x17, 0x2A))
        for p in article.find_all("p", recursive=True):
            text = visible_text(p)
            if not text:
                continue
            self.add_plain(text, indent=0, size=11)
        ol = article.find("ol")
        if ol:
            self.add_list(ol, ordered=True)

    def add_callout(self, box: Tag):
        title = box.find("p")
        if title:
            self.add_plain(visible_text(title), bold=True, indent=0, size=12, space_before=8, space_after=2)
        body = box.find("div")
        if body:
            for child in body.children:
                if isinstance(child, Tag) and child.name == "p":
                    self.add_body(child, indent=0)
                elif isinstance(child, Tag):
                    self.walk(child)
        elif title is None:
            self.add_plain(visible_text(box), indent=0)

    def add_codeblock(self, pre: Tag):
        text = pre.get_text("\n")
        p = self.doc.add_paragraph()
        set_paragraph_format(p, first_indent=0, space_before=4, space_after=8, line=1.15)
        run = p.add_run(text.strip("\n"))
        set_run_font(run, ascii_font=EN_CODE, east=CN_BODY, size=9.5)

    def walk(self, node: Tag):
        for child in node.children:
            if not isinstance(child, Tag):
                continue
            cls = classes(child)
            if "print:hidden" in cls:
                continue
            name = child.name

            if name == "h1":
                self.add_heading_text(visible_text(child), 0)
                continue
            if name == "p" and "tracking-[0.18em]" in cls:
                self.pending_kicker = visible_text(child)
                continue
            if name == "h2":
                title = visible_text(child)
                if self.pending_kicker:
                    title = f"{self.pending_kicker}　{title}"
                    self.pending_kicker = ""
                self.add_heading_text(title, 1)
                continue
            if name == "h3":
                self.add_heading_text(visible_text(child), 2)
                continue
            if name == "p" and "font-serif" in cls:
                self.add_body(child, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
                continue
            if name == "p":
                if not visible_text(child):
                    continue
                self.add_body(child)
                continue
            if name == "ul":
                self.add_list(child, ordered=False)
                continue
            if name == "ol":
                self.add_list(child, ordered=True)
                continue
            if name == "table":
                self.add_table(child)
                continue
            if name == "figure":
                self.add_figure(child)
                continue
            if name == "pre":
                self.add_codeblock(child)
                continue
            if name == "article" and "border-amber-200/80" in cls:
                self.add_qa(child)
                continue
            if name == "article" and child.find("header"):
                self.add_algo(child)
                continue
            if name == "div" and "rounded-lg border" in cls and "px-4 py-3" in cls:
                self.add_callout(child)
                continue
            self.walk(child)

    def build(self) -> bytes:
        article = self.soup.find("article")
        if not article:
            raise SystemExit("print page has no <article>")
        self.walk(article)
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def fetch_html() -> str:
    if len(sys.argv) > 1:
        return Path(sys.argv[1]).read_text(encoding="utf-8")
    with urllib.request.urlopen(PRINT_URL, timeout=30) as resp:
        return resp.read().decode("utf-8", errors="replace")


def main() -> None:
    html = fetch_html()
    data = Builder(html).build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(data)
    docs_copy = ROOT / "docs" / "继电器接点寿命测试软件-课题报告.docx"
    docs_copy.write_bytes(data)
    print(f"wrote {OUT} ({len(data)} bytes)")
    print(f"wrote {docs_copy}")


if __name__ == "__main__":
    main()
